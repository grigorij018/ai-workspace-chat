import io
import hashlib
import re
import time
import uuid
from typing import Iterable

from pptx import Presentation
from pptx.util import Inches, Pt

from open_webui.models.files import FileForm, FileModel, Files
from open_webui.storage.provider import Storage


PPTX_CONTENT_TYPE = 'application/vnd.openxmlformats-officedocument.presentationml.presentation'


def _message_text(message: dict) -> str:
    content = message.get('content', '')
    if isinstance(content, str):
        return content.strip()
    if isinstance(content, list):
        return '\n'.join(part.get('text', '') for part in content if part.get('type') == 'text').strip()
    return ''


def summarize_chat_for_pptx(messages: Iterable[dict], max_chars: int = 6000) -> str:
    lines = []
    for message in messages:
        role = message.get('role', 'message')
        text = _message_text(message)
        if text:
            lines.append(f'{role}: {text}')

    transcript = '\n\n'.join(lines).strip()
    return transcript[-max_chars:] if len(transcript) > max_chars else transcript


def _clean_bullet(text: str) -> str:
    text = re.sub(r'\s+', ' ', text).strip()
    return text[:180].rstrip()


def _extract_bullets(transcript: str, fallback: list[str], limit: int = 4) -> list[str]:
    candidates = []
    for line in transcript.splitlines():
        line = re.sub(r'^(user|assistant|system):\s*', '', line.strip(), flags=re.I)
        if not line:
            continue
        if any(marker in line.lower() for marker in ('file', 'url', 'research', 'rag', 'chat', 'memory', 'demo')):
            candidates.append(_clean_bullet(line))

    seen = set()
    bullets = []
    for candidate in candidates:
        key = candidate.lower()
        if candidate and key not in seen:
            seen.add(key)
            bullets.append(candidate)
        if len(bullets) >= limit:
            break

    return bullets or fallback


def build_chat_pptx(transcript: str) -> bytes:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    sections = [
        ('AI Workspace', ['Unified chat for files, URLs, research, memory, and presentation output.']),
        ('Problem', ['Workflows are split across tools.', 'Sources, files, and memory are hard to keep in one place.']),
        ('Solution', _extract_bullets(transcript, ['One chat flow routes files, links, search, and generation.'])),
        ('Architecture', ['OpenWebUI chat UX', 'MWS GPT routing', 'RAG parsing, chunking, embeddings, retrieval', 'Transparent memory and citations']),
        ('Demo Flow', ['Attach a file', 'Paste a URL', 'Enable research', 'Ask for a PPTX from the same chat']),
        ('Value', ['Fewer context switches', 'Cited answers', 'Reusable presentation artifact']),
        ('Roadmap', ['Credentialed MWS smoke run', 'More source connectors', 'Presentation template controls']),
    ]

    for index, (title, bullets) in enumerate(sections):
        layout = prs.slide_layouts[0] if index == 0 else prs.slide_layouts[1]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = title
        placeholder = slide.placeholders[1]
        placeholder.text = ''
        text_frame = placeholder.text_frame
        text_frame.clear()
        for bullet in bullets[:5]:
            paragraph = text_frame.add_paragraph()
            paragraph.text = bullet
            paragraph.level = 0
            paragraph.font.size = Pt(22 if index else 24)

    output = io.BytesIO()
    prs.save(output)
    return output.getvalue()


def create_chat_pptx_file(user_id: str, messages: list[dict]) -> FileModel:
    transcript = summarize_chat_for_pptx(messages)
    pptx_bytes = build_chat_pptx(transcript)
    file_id = str(uuid.uuid4())
    name = f'ai-workspace-research-{int(time.time())}.pptx'
    stored_name = f'{file_id}_{name}'

    _, file_path = Storage.upload_file(
        io.BytesIO(pptx_bytes),
        stored_name,
        {
            'OpenWebUI-User-Id': user_id,
            'OpenWebUI-File-Id': file_id,
        },
    )

    return Files.insert_new_file(
        user_id,
        FileForm(
            id=file_id,
            filename=name,
            path=file_path,
            hash=hashlib.sha256(pptx_bytes).hexdigest(),
            data={'status': 'completed', 'content': transcript},
            meta={
                'name': name,
                'content_type': PPTX_CONTENT_TYPE,
                'size': len(pptx_bytes),
                'generated_by': 'pptx_generation',
            },
        ),
    )
